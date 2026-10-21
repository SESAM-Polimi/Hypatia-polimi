from hypatia.backend.ModelSettings import ModelSettings
from hypatia.backend.ModelData import ModelData
from hypatia.utility.constants import ModelMode
from hypatia.utility.constants import OptimizationMode
from hypatia.utility.constants import EnsureFeasibility
from hypatia.utility.utility import CAPACITY_GROUP_SHEETS
from openpyxl import load_workbook
import pandas as pd
import cvxpy as cp
import os
import shutil
import warnings
from hypatia.error_log.Exceptions import (
    WrongInputMode,
    DataNotImported,
    ResultOverWrite,
    SolverNotFound,
)


def _read_optional_group_sheet(file_path: str, sheet_name: str) -> pd.DataFrame:
    sheet = pd.read_excel(file_path, sheet_name=sheet_name)
    return sheet.loc[:, ~sheet.columns.astype(str).str.startswith("Unnamed:")]


def _write_empty_capacity_group_sheets(writer):
    for value in CAPACITY_GROUP_SHEETS.values():
        pd.DataFrame(columns=value["columns"]).to_excel(
            writer,
            sheet_name=value["sheet_name"],
            index=False,
        )


def _option_sheets(settings):
    """{file name: {sheet name: (option, template entry)}} for the optional
    constraints that are not enabled: the sheets they would read if they were."""
    from hypatia.backend.constraints.ConstraintList import OPTIONAL_CONSTRAINTS

    sheets = {}
    for option, cls in OPTIONAL_CONSTRAINTS.items():
        if option in settings.optional_constraints:
            continue
        for key, entry in (cls.required_global_parameters(cls, settings) or {}).items():
            sheets.setdefault("parameters_global.xlsx", {})[entry["sheet_name"]] = (option, entry)
        for key, entry in (cls.required_trade_parameters(cls, settings) or {}).items():
            sheets.setdefault("parameters_connections.xlsx", {})[entry["sheet_name"]] = (option, entry)
        regional = cls.required_regional_parameters(cls, settings) or {}
        for region, entries in regional.items():
            for key, entry in entries.items():
                sheets.setdefault(f"parameters_{region}.xlsx", {})[entry["sheet_name"]] = (option, entry)
    return sheets


def _holds_data(file_path, sheet, entry):
    """True if the sheet holds values other than the option's default.

    Sheets left at their defaults (e.g. ramp rates of 1, "no limit") are what
    create_data_excels writes when an option was on once; ignoring them changes
    nothing, so they are not worth a warning. A sheet that cannot be read in the
    expected shape counts as holding data.
    """
    try:
        frame = pd.read_excel(
            file_path,
            sheet_name=sheet,
            index_col=list(range(entry["index"].nlevels)),
            header=list(range(entry["columns"].nlevels)),
        )
    except Exception:
        return True
    default = entry["value"]
    return bool((frame.astype(object) != default).to_numpy().any())


def _warn_ignored_option_sheets(settings, file_path, sheet_names, read_sheets):
    """Warn about filled sheets of a disabled optional constraint in a file.

    Without this the sheet is silently ignored: a model run without
    chp_operating_range ignores a filled CHP sheet and gives different results.
    """
    file_name = os.path.basename(file_path)
    for sheet, (option, entry) in _option_sheets(settings).get(file_name, {}).items():
        if sheet in sheet_names and sheet not in read_sheets and _holds_data(file_path, sheet, entry):
            warnings.warn(
                f"{file_name} contains the sheet '{sheet}', filled with data, which is read "
                f"only with the optional constraint '{option}'. That option is not enabled, so "
                f"the sheet is ignored. To use it: Model(..., optional_constraints=['{option}']).",
                UserWarning,
                stacklevel=3,
            )


def read_settings(path: str, mode: ModelMode, optimization: OptimizationMode, ensure_feasibility: EnsureFeasibility, optional_constraints=()) -> ModelSettings:
    wb_glob = load_workbook(r"{}/global.xlsx".format(path))
    sets_glob = wb_glob["Sets"]

    global_settings = {}
    for entry, data_boundary in sets_glob.tables.items():
        data_glob = sets_glob[data_boundary]
        content = [[cell.value for cell in ent] for ent in data_glob]
        header = content[0]
        rest = content[1:]
        df = pd.DataFrame(rest, columns=header)
        global_settings[entry] = df

    regional_settings = {}
    for reg in list(global_settings["Regions"]["Region"]):
        wb = load_workbook(r"{}/{}.xlsx".format(path, reg))
        sets = wb["Sets"]
        settings = {}
        for entry, data_boundary in sets.tables.items():
            data = sets[data_boundary]
            content = [[cell.value for cell in ent] for ent in data]
            header = content[0]
            rest = content[1:]
            df = pd.DataFrame(rest, columns=header)
            settings[entry] = df

        regional_settings[reg] = settings

    return ModelSettings(
        mode, optimization, ensure_feasibility, global_settings, regional_settings,
        optional_constraints=optional_constraints,
    )


def write_parameters_files(settings: ModelSettings, path: str, force_rewrite: bool = False):
    if os.path.exists(path):
        if not force_rewrite:
            raise ResultOverWrite(
                f"Folder {path} already exists. To over write"
                f" the parameter files, use force_rewrite=True."
            )
        else:
            shutil.rmtree(path)
    os.mkdir(path)
    if settings.multi_node:
        with pd.ExcelWriter(
            r"{}/parameters_connections.xlsx".format(path)
        ) as writer:
            for key, value in settings.trade_parameters_template.items():
                connection_data = pd.DataFrame(
                    value["value"], index=value["index"], columns=value["columns"]
                )
                connection_data.to_excel(writer, sheet_name=value["sheet_name"])

        with pd.ExcelWriter(r"{}/parameters_global.xlsx".format(path)) as writer:
            for key, value in settings.global_parameters_template.items():
                global_data = pd.DataFrame(
                    value["value"], index=value["index"], columns=value["columns"]
                )
                global_data.to_excel(writer, sheet_name=value["sheet_name"])
            _write_empty_capacity_group_sheets(writer)

    for reg in settings.regions:
        with pd.ExcelWriter(r"{}/parameters_{}.xlsx".format(path, reg)) as writer:
            for key, value in settings.regional_parameters_template[reg].items():
                regional_data = pd.DataFrame(
                    value["value"], index=value["index"], columns=value["columns"]
                )
                regional_data.to_excel(writer, sheet_name=value["sheet_name"])
            if not settings.multi_node:
                _write_empty_capacity_group_sheets(writer)

def read_parameters(settings: ModelSettings, path: str) -> ModelData:
    global_parameters = None
    if settings.global_parameters_template != None:
        global_parameters = {}
        global_file = r"{}/parameters_global.xlsx".format(path)
        global_excel = pd.ExcelFile(global_file)
        for key, value in settings.global_parameters_template.items():
            global_parameters[key] = pd.read_excel(
                global_file,
                sheet_name=value["sheet_name"],
                index_col=list(range(0, value["index"].nlevels)),
                header=list(range(0, value["columns"].nlevels)),
            )

        _warn_ignored_option_sheets(
            settings, global_file, global_excel.sheet_names,
            {value["sheet_name"] for value in settings.global_parameters_template.values()},
        )
        for key, value in CAPACITY_GROUP_SHEETS.items():
            if value["sheet_name"] in global_excel.sheet_names:
                global_parameters[key] = _read_optional_group_sheet(
                    global_file,
                    value["sheet_name"],
                )

    trade_parameters = None
    if settings.trade_parameters_template != None:
        trade_parameters = {}
        trade_file = r"{}/parameters_connections.xlsx".format(path)
        _warn_ignored_option_sheets(
            settings, trade_file, pd.ExcelFile(trade_file).sheet_names,
            {value["sheet_name"] for value in settings.trade_parameters_template.values()},
        )
        for key, value in settings.trade_parameters_template.items():
            trade_parameters[key] = pd.read_excel(
                r"{}/parameters_connections.xlsx".format(path),
                sheet_name=value["sheet_name"],
                index_col=list(range(0, value["index"].nlevels)),
                header=list(range(0, value["columns"].nlevels)),
            )

    regional_parameters = {}
    for region in settings.regions:
        parameters = {}
        regional_file = r"{}/parameters_{}.xlsx".format(path, region)
        regional_excel = pd.ExcelFile(regional_file)
        for key, value in settings.regional_parameters_template[region].items():
            parameters[key] = pd.read_excel(
                regional_file,
                sheet_name=value["sheet_name"],
                index_col=list(range(0, value["index"].nlevels)),
                header=list(range(0, value["columns"].nlevels)),
            )

        _warn_ignored_option_sheets(
            settings, regional_file, regional_excel.sheet_names,
            {value["sheet_name"] for value in settings.regional_parameters_template[region].values()},
        )
        for key, value in CAPACITY_GROUP_SHEETS.items():
            if value["sheet_name"] in regional_excel.sheet_names:
                parameters[key] = _read_optional_group_sheet(
                    regional_file,
                    value["sheet_name"],
                )
        regional_parameters[region] = parameters

    return ModelData(
        settings,
        global_parameters,
        trade_parameters,
        regional_parameters,
    )
